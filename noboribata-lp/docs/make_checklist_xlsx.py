# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

JP = "Meiryo"
NAVY = "00406F"; BLUE = "0066CC"; MIST = "EEF4FA"; LINE = "D7E2EE"
YELLOW = "FFF6C2"; GREY = "666666"

thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

def F(sz=10, b=False, color="16202C"):
    return Font(name=JP, size=sz, bold=b, color=color)

wb = Workbook()

# ============================================================
# シート1 はじめに
# ============================================================
ws = wb.active
ws.title = "はじめに"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 92


def head(r, text):
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name=JP, size=11, bold=True, color="FFFFFF")
    c.alignment = Alignment(vertical="center", indent=1)
    for cc in range(2, 4):
        ws.cell(row=r, column=cc).fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 22

def para(r, text, size=10, color="16202C"):
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name=JP, size=size, color=color)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c.alignment = Alignment(vertical="center", wrap_text=True)

def pair(r, k, v):
    a = ws.cell(row=r, column=2, value=k)
    a.font = F(10, True, NAVY)
    a.alignment = Alignment(vertical="center")
    b = ws.cell(row=r, column=3, value=v)
    b.font = F(10)
    b.alignment = Alignment(vertical="center", wrap_text=True)

ws["B2"] = "のぼり旗の外注パートナー（新ページ）ご確認事項シート"
ws["B2"].font = Font(name=JP, size=16, bold=True, color=NAVY)
ws.merge_cells("B2:C2"); ws.row_dimensions[2].height = 26
ws["B3"] = "株式会社後藤染工 御中"
ws["B3"].font = Font(name=JP, size=11, color=GREY)
ws.merge_cells("B3:C3")

head(5, "このシートについて")
para(6, "新しいWebページを正式に公開するにあたり、御社に確認させていただきたいことをまとめました。")
para(7, "黄色いセルにご記入のうえ、このファイルをそのままご返信ください。")
para(8, "すべて埋まっていなくても構いません。分かる範囲でご記入いただければ、残りは改めてご相談させてください。")

head(10, "シートの構成")
pair(11, "① はじめに", "このページです。記入のしかたが書いてあります。")
pair(12, "② ご確認事項", "ご確認したいこと27項目の一覧です。E列（黄色）にご回答をご記入ください。")
pair(13, "③ 納期・送料・支払条件", "表の形でご記入いただく項目です。黄色いセルを埋めてください。")

head(15, "記入のしかた")
pair(16, "黄色のセル", "ご記入いただく欄です。ここだけ埋めてください。")
pair(17, "白・灰色のセル", "こちらで記入済みの説明文です。変更は不要です。")
pair(18, "オレンジ色の行", "書き方の見本です。消していただいて構いません。")
pair(19, "優先度「高」", "ページ公開前に必要な項目です。優先してご記入ください。")
pair(20, "優先度「中」", "公開後の追記でも問題ありませんが、あるとAI検索に引用されやすくなります。")
pair(21, "分からない項目", "「未定」「相談したい」とご記入いただければ結構です。空欄のままでも構いません。")

head(23, "ご記入状況（自動で計算されます）")
a = ws.cell(row=24, column=2, value="ご記入済みの項目数"); a.font = F(10, True, NAVY)
b = ws.cell(row=24, column=3, value='=COUNTA(\'ご確認事項\'!E6:E32)&" ／ 27 項目"'); b.font = F(11, True)
a2 = ws.cell(row=25, column=2, value="うち優先度「高」"); a2.font = F(10, True, NAVY)
b2 = ws.cell(row=25, column=3, value='=COUNTIFS(\'ご確認事項\'!G6:G32,"高",\'ご確認事項\'!E6:E32,"<>")&" ／ "&COUNTIF(\'ご確認事項\'!G6:G32,"高")&" 項目"')
b2.font = F(11, True)

head(27, "参考：なぜこれらを伺うのか")
para(28, "※「AEO（エーイーオー）」＝ ChatGPTなどのAIが質問に答えるとき、その回答の材料として自社サイトを引用してもらうための対策のことです。", 9, GREY)
para(29, "※ AIは画像の中の文字が読めません。数字や条件を「文章」として書いておくことが、引用されるための条件になります。", 9, GREY)
para(30, "※ このシートでお伺いしている項目は、いずれも「のぼり旗を外注したい会社」がAIに質問しやすい内容です。", 9, GREY)
para(31, "※ 納期・送料・支払条件は現在のサイトに記載がないため、こちらで推測して書くことはしていません。", 9, GREY)
for rr in (6,7,8,11,12,13,16,17,18,19,20,21,28,29,30,31):
    ws.row_dimensions[rr].height = 18

# ============================================================
# シート2 ご確認事項
# ============================================================
ws2 = wb.create_sheet("ご確認事項")
ws2.sheet_view.showGridLines = False
widths = {"A": 5, "B": 16, "C": 40, "D": 46, "E": 44, "F": 34, "G": 8}
for k, v in widths.items():
    ws2.column_dimensions[k].width = v

ws2["A1"] = "ご確認事項一覧"
ws2["A1"].font = Font(name=JP, size=15, bold=True, color=NAVY)
ws2.merge_cells("A1:G1")
ws2.row_dimensions[1].height = 24
ws2["A2"] = "黄色のE列にご回答をご記入ください。分かる範囲で結構です。"
ws2["A2"].font = F(10, False, GREY)
ws2.merge_cells("A2:G2")
ws2["A3"] = "記入例（この行は書き方の見本です。消していただいて構いません）"
ws2["A3"].font = F(9, True, "B06000")
ws2.merge_cells("A3:G3")

HEAD = ["No", "分類", "ご確認したいこと", "なぜ必要か（ページでの使いみち）", "ご回答（←ご記入ください）", "記入例・選択肢", "優先度"]
for i, h in enumerate(HEAD, start=1):
    c = ws2.cell(row=5, column=i, value=h)
    c.font = Font(name=JP, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box
ws2.row_dimensions[5].height = 32

DATA = [
 ("公開・運用","本番ドメイン（www.noboribata-gaityuu-partner.com）へ新ページを切り替えてよいでしょうか。希望時期もご記入ください。","切り替え日が決まらないと公開作業に入れません。","例）2026年11月10日に切り替え希望","高"),
 ("公開・運用","現在のサイトのサーバー・ドメインを管理している会社名とご担当者を教えてください。","切り替え作業の連絡先になります。","例）○○ホスティング株式会社 ／ 担当 □□様","高"),
 ("公開・運用","お問い合わせフォームの送信方式はどちらにしますか。A：お客様のメールソフトが起動する方式（追加費用なし）／B：現サイトと同じ確認画面つきフォーム（サーバー側の設定が別途必要）","Bの場合は別途お見積りが必要になります。","A または B","高"),
 ("公開・運用","フォームの受信先は info@gotosenko.co.jp で問題ないでしょうか。別アドレスの場合はご記入ください。","問い合わせの届け先を間違えないためです。","例）info@gotosenko.co.jp でよい","高"),
 ("公開・運用","電話番号 0568-35-6651 をページに掲載して問題ないでしょうか。また電話の受付時間と定休日を教えてください。","「何時まで電話できるか」はAIによく聞かれる項目です。","例）掲載可。平日8:30〜17:30、土日祝休","高"),
 ("公開・運用","Google Search Console・Googleアナリティクス（アクセス数を見る無料ツール）のアカウントはお持ちですか。","公開後の効果測定に使います。無ければこちらで作成します。","例）持っていない／こちらで作ってほしい","中"),
 ("未掲載の情報","納期の目安を教えてください（枚数別）。※シート③にご記入ください","「何日で届くか」はAIが最も多く聞かれる質問のひとつです。","シート③「納期・送料・支払条件」へ","高"),
 ("未掲載の情報","送料・配送条件を教えてください。※シート③にご記入ください","送料の有無は発注判断を左右するため、明記すると引用されやすくなります。","シート③「納期・送料・支払条件」へ","高"),
 ("未掲載の情報","入稿データの締切時間を教えてください（何時までの入稿でその日のうちに処理できますか）。","「今日中に入稿すれば間に合うか」の判断材料になります。","例）平日15:00までの入稿で当日受付","中"),
 ("未掲載の情報","最小ロットは「1枚から」で間違いないでしょうか。サイズや素材によって異なる場合はご記入ください。","現サイトのQ&Aに記載がありますが、念のための確認です。","例）1枚から可。ただし防炎品は10枚から","高"),
 ("未掲載の情報","お支払い条件を教えてください（初回・2回目以降）。※シート③にご記入ください","取引を検討する法人がほぼ必ず確認する項目です。","シート③「納期・送料・支払条件」へ","高"),
 ("未掲載の情報","印刷サンプル・色校正は無料ですか、有料ですか。有料の場合は費用も教えてください。","現サイトには「対応可能」とだけ書かれており、費用の記載がありません。","例）基本メディアのサンプルは無料。本紙校正は1枚○○円","中"),
 ("未掲載の情報","入稿テンプレートの配布ページのURLがあれば教えてください。","ページからリンクすると、問い合わせ前の準備が進みます。","例）https://www.gotosenko.co.jp/template/","中"),
 ("未掲載の情報","対応できる縫製・加工の種類を教えてください（チチ、袋縫い、三方三巻、ヒートカットなど）。","「この加工はできますか」という質問への答えになります。","例）チチ／袋／三方三巻／棒袋 に対応","中"),
 ("未掲載の情報","防炎ラベル（防炎性能の認定表示）の発行は可能ですか。","商業施設やイベント向けの案件で必ず確認される項目です。","例）防炎ポンジ使用時はラベル発行可（1枚○○円）","中"),
 ("未掲載の情報","のぼり旗以外の製品（のれん・はっぴ・タオルなど）の最小ロットを教えてください。","製品ごとに条件が違う場合、書いておくと問い合わせが減ります。","例）タオルは100枚から、はっぴは10枚から","中"),
 ("未掲載の情報","印刷不良・誤納品があった場合の対応方針を教えてください。","「もし失敗したらどうなるか」は外注先選びの決め手になります。","例）弊社起因の不良は無償で再製作","中"),
 ("未掲載の情報","デザイン制作やデータ修正の代行は可能ですか。可能な場合の費用も教えてください。","Aiデータを用意できない会社からの相談が拾えるようになります。","例）簡単な文字修正は無償、新規デザインは○○円から","中"),
 ("未掲載の情報","代理発送の条件を教えてください（送り状の差出人名義を御社名にできるか、同梱物を入れられるか）。","代理発送は御社の強みですが、条件が書かれていません。","例）送り状は依頼主名義に変更可。チラシ同梱も可","中"),
 ("未掲載の情報","価格表は現在「2025年12月時点」と記載しています。最新の価格に更新する必要はありますか。","古い日付のままだとAIにも利用者にも不安を与えます。","例）そのままでよい／2026年○月時点に更新したい","高"),
 ("掲載内容の確認","「業界最安値水準」「24時間プリント」「年間実績100万枚超」の表現をそのまま使ってよいでしょうか。根拠があれば併せて教えてください。","根拠を1行添えると信頼性が上がり、AIにも引用されやすくなります。","例）そのまま使用可。年間出荷枚数は社内集計値","高"),
 ("掲載内容の確認","取引実績のある業種や企業名を掲載してよいでしょうか。掲載可能なものがあれば教えてください。","実名や業種が出ると、AIが「実績のある会社」と判断しやすくなります。","例）業種のみ可（自動車販売、パチンコ、不動産 など）","中"),
 ("掲載内容の確認","工場・設備・作業風景・社員の写真をご提供いただけますか。","写真があると人にもAIにも「実在する工場」と伝わります。","例）工場外観と印刷機の写真を提供可","中"),
 ("掲載内容の確認","ISOなどの認証、受賞歴、加盟団体があれば教えてください。","第三者からの評価はAI検索で重視されます。","例）特になし／○○工業会 加盟","中"),
 ("掲載内容の確認","プライバシーポリシー・特定商取引法に基づく表記のページは新ページにも置きますか。","フォームで個人情報を受け取るため、通常は設置を推奨します。","例）コーポレートサイトの既存ページにリンクしたい","高"),
 ("掲載内容の確認","WEB会議で使用しているツールを教えてください（Zoom／Teams／Google Meet など）。","「WEB会議も可」と書くなら、ツール名まであると親切です。","例）Zoom と Teams に対応","中"),
 ("その他","そのほか、新しいページに載せたい情報・載せたくない情報があればご記入ください。","","例）繁忙期（3月・9月）は納期が延びる旨を書いてほしい","中"),
]

yfill = PatternFill("solid", fgColor=YELLOW)
center = Alignment(horizontal="center", vertical="center")
wrap = Alignment(vertical="center", wrap_text=True)

# 記入例の行
ex = ["例", "公開・運用", "（書き方の見本）本番ドメインへ切り替えてよいですか。", "切り替え日が決まらないと作業に入れません。",
      "2026年11月10日に切り替え希望。作業は平日午前でお願いします。", "例）2026年11月10日に切り替え希望", "高"]
for i, v in enumerate(ex, start=1):
    c = ws2.cell(row=4, column=i, value=v)
    c.font = Font(name=JP, size=9, italic=True, color="B06000")
    c.alignment = center if i in (1, 7) else wrap
    c.border = box
    c.fill = PatternFill("solid", fgColor="FFF9E6")
ws2.row_dimensions[4].height = 40

r = 6
for n, (cat, q, why, ex_, pri) in enumerate(DATA, start=1):
    vals = [n, cat, q, why, None, ex_, pri]
    for i, v in enumerate(vals, start=1):
        c = ws2.cell(row=r, column=i, value=v)
        c.border = box
        if i in (1, 7):
            c.font = F(10, True) if i == 7 else F(10)
            c.alignment = center
        elif i == 2:
            c.font = F(9, True, NAVY)
            c.alignment = wrap
            c.fill = PatternFill("solid", fgColor=MIST)
        elif i == 5:
            c.fill = yfill
            c.font = F(10)
            c.alignment = wrap
        elif i == 6:
            c.font = F(9, False, GREY)
            c.alignment = wrap
        else:
            c.font = F(10)
            c.alignment = wrap
        if i == 7 and v == "高":
            c.font = Font(name=JP, size=10, bold=True, color="D41E00")
    def lines(text, width_chars):
        if not text:
            return 1
        n = 0
        for ch in text:
            n += 1 if ord(ch) < 0x2000 else 2
        return max(1, -(-n // max(1, int(width_chars) - 4)))
    need = max(lines(q, 40), lines(why, 46), lines(ex_, 34))
    ws2.row_dimensions[r].height = max(34, need * 16 + 12)
    r += 1

ws2.freeze_panes = "A6"
ws2.auto_filter.ref = f"A5:G{r-1}"

# ============================================================
# シート3 納期・送料・支払条件
# ============================================================
ws3 = wb.create_sheet("納期・送料・支払条件")
ws3.sheet_view.showGridLines = False
for k, v in {"A": 4, "B": 22, "C": 24, "D": 24, "E": 24, "F": 40}.items():
    ws3.column_dimensions[k].width = v

def title(row, text):
    c = ws3.cell(row=row, column=2, value=text)
    c.font = Font(name=JP, size=12, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(vertical="center", indent=1)
    for cc in range(2, 7):
        ws3.cell(row=row, column=cc).fill = PatternFill("solid", fgColor=BLUE)
    ws3.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    ws3.row_dimensions[row].height = 22

def note(row, text):
    c = ws3.cell(row=row, column=2, value=text)
    c.font = F(9, False, GREY)
    ws3.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    c.alignment = Alignment(vertical="center", wrap_text=True)

def table(row, headers, rowlabels, input_cols, example=None):
    for i, h in enumerate(headers):
        c = ws3.cell(row=row, column=2 + i, value=h)
        c.font = Font(name=JP, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    ws3.row_dimensions[row].height = 30
    rr = row + 1
    if example:
        for i, v in enumerate(example):
            c = ws3.cell(row=rr, column=2 + i, value=v)
            c.font = Font(name=JP, size=9, italic=True, color="B06000")
            c.fill = PatternFill("solid", fgColor="FFF9E6")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = box
        ws3.row_dimensions[rr].height = 24
        rr += 1
    for lab in rowlabels:
        c = ws3.cell(row=rr, column=2, value=lab)
        c.font = F(10, True, NAVY)
        c.fill = PatternFill("solid", fgColor=MIST)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
        for i in input_cols:
            cc = ws3.cell(row=rr, column=2 + i)
            cc.fill = yfill
            cc.border = box
            cc.font = F(10)
            cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws3.row_dimensions[rr].height = 24
        rr += 1
    return rr

ws3["B1"] = "納期・送料・お支払い条件 ご記入シート"
ws3["B1"].font = Font(name=JP, size=15, bold=True, color=NAVY)
ws3.merge_cells("B1:F1")
ws3.row_dimensions[1].height = 24
ws3["B2"] = "黄色のセルにご記入ください。分かる範囲で結構です。オレンジ色の行は書き方の見本です。"
ws3["B2"].font = F(10, False, GREY)
ws3.merge_cells("B2:F2")

title(4, "① 納期の目安（450mm×1,500mm・ポンジ・チチ付きの場合）")
r = table(5, ["ご注文枚数", "通常納期（営業日）", "特急対応の可否", "特急時の納期", "備考"],
          ["10枚", "50枚", "100枚", "500枚", "1000枚"], [1, 2, 3, 4],
          example=["（例）10枚", "5営業日", "可（割増あり）", "3営業日", "入稿データ確定日からの日数"])
note(r + 1, "※「営業日」＝土日祝を除いた日数です。データ確定日から出荷日までで数えていただくと分かりやすいです。")
note(r + 2, "※ サイズや素材で納期が変わる場合は、備考欄にご記入ください。")

title(r + 4, "② 送料・配送条件")
r2 = table(r + 5, ["お届け先", "送料（円・税抜）", "送料無料になる条件", "お届け日数の目安", "備考"],
           ["本州", "北海道", "四国", "九州", "沖縄・離島"], [1, 2, 3, 4],
           example=["（例）本州", "1,200", "10,000円以上で無料", "出荷翌日", "1梱包あたり"])
note(r2 + 1, "※ 代理発送（使用先への直送）の場合に追加費用が発生するかどうかも、備考欄にご記入ください。")

title(r2 + 3, "③ お支払い条件")
r3 = table(r2 + 4, ["区分", "お支払い方法", "締め日", "お支払い期日", "備考"],
           ["初回のお取引", "2回目以降", "個人・個人事業主"], [1, 2, 3, 4],
           example=["（例）初回のお取引", "銀行振込（前金）", "—", "出荷前までにご入金", "請求書を先に発行"])
note(r3 + 1, "※ 掛け取引（後払い）に与信審査が必要な場合は、備考欄にその旨をご記入ください。")

title(r3 + 3, "④ 繁忙期・休業期間")
r4 = table(r3 + 4, ["時期", "納期への影響", "休業の有無", "期間", "備考"],
           ["年末年始", "ゴールデンウィーク", "お盆", "その他の繁忙期"], [1, 2, 3, 4],
           example=["（例）年末年始", "通常より+3営業日", "休業あり", "12/29〜1/4", "12/20以降の入稿は年明け出荷"])
note(r4 + 1, "※ 休業期間を書いておくと「今頼んで間に合うか」の問い合わせが減り、AIの回答にも反映されます。")

for sht in wb.worksheets:
    sht.sheet_properties.tabColor = BLUE
    sht.page_setup.orientation = "landscape"
    sht.page_setup.fitToWidth = 1
    sht.page_setup.fitToHeight = 0
    sht.sheet_properties.pageSetUpPr.fitToPage = True
    sht.print_options.horizontalCentered = True
ws2.print_title_rows = "5:5" 

out = "/home/user/TEST/noboribata-lp/docs/後藤染工_ご確認事項シート_20261027.xlsx"
wb.save(out)
print("saved", out)
